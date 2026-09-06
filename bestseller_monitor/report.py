"""差分计算与 CSV/Excel 导出。"""
from __future__ import annotations

import csv
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from .config import Config
from .db import Database, CST

log = logging.getLogger(__name__)


def _key(row) -> tuple[str, str, str]:
    return (row["shop_key"], row["offer_id"], row["sku_id"])


def update_stock_deltas(db: Database, round_id: int) -> None:
    """用最近一个已完成轮次给本轮成功快照回填 stock_delta。"""
    prev = db.previous_complete_round(round_id)
    if prev is None:
        log.info("没有可对比的上一轮，本轮不计算库存变化。")
        return
    prev_rows = db.success_rows(int(prev["id"]))
    prev_map = {_key(r): r for r in prev_rows}
    cur_rows = db.success_rows(round_id)
    for row in cur_rows:
        p = prev_map.get(_key(row))
        if p is None:
            db.update_delta(int(row["id"]), 0)
            continue
        cur_stock = row["sku_stock"]
        prev_stock = p["sku_stock"]
        if cur_stock is None or prev_stock is None:
            db.update_delta(int(row["id"]), 0)
            continue
        db.update_delta(int(row["id"]), int(cur_stock) - int(prev_stock))
    db.commit()


def export_round_csv(db: Database, cfg: Config, round_id: int) -> list[Path]:
    """导出 shop_offers 与 snapshots 两个 CSV。"""
    paths: list[Path] = []
    stamp = _stamp(db, round_id)

    off_path = cfg.data_dir / f"shop_offers_{stamp}.csv"
    rows = db.conn.execute(
        "SELECT shop_key, shop_name, rank, offer_id, product_url, list_title, list_price "
        "FROM shop_offers WHERE round_id=? ORDER BY shop_key, rank",
        (round_id,),
    ).fetchall()
    with open(off_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["店铺", "排名", "offer_id", "商品链接", "列表标题", "列表价格"])
        for r in rows:
            w.writerow([r["shop_name"], r["rank"], r["offer_id"], r["product_url"],
                        r["list_title"] or "", r["list_price"] or ""])
    paths.append(off_path)

    snap_path = cfg.data_dir / f"snapshots_{stamp}.csv"
    cols = [
        "shop_name", "product_url", "product_name", "sku_name", "sku_price",
        "sku_stock", "stock_delta", "collected_at",
    ]
    rows = db.success_rows(round_id)
    with open(snap_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["店铺", "商品链接", "商品名称", "SKU名称", "SKU价格", "SKU库存量", "库存变化", "采集时间"])
        for r in rows:
            w.writerow([r[cols[0]], r["product_url"], r["product_name"], r["sku_name"],
                        r["sku_price"], r["sku_stock"], r["stock_delta"], r["collected_at"]])
    paths.append(snap_path)
    log.info("CSV 已导出：%s", ", ".join(str(p) for p in paths))
    return paths


def export_round_excel(db: Database, cfg: Config, round_id: int) -> Path:
    """导出 Excel 日报（openpyxl）。"""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请执行 pip install -r requirements.txt") from exc

    stamp = _stamp(db, round_id)
    wb = Workbook()
    _sheet_榜单(wb, db, round_id)
    _sheet_差分(wb, db, round_id)
    _sheet_汇总(wb, db, round_id)
    _sheet_补货(wb, db, round_id)
    _sheet_榜单变化(wb, db, round_id)
    _sheet_失败(wb, db, round_id)
    out = cfg.output_dir / f"日报_{stamp}.xlsx"
    wb.save(out)
    log.info("Excel 已导出：%s", out)
    return out


def _stamp(db: Database, round_id: int) -> str:
    row = db.conn.execute("SELECT started_at FROM rounds WHERE id=?", (round_id,)).fetchone()
    if not row:
        return datetime.now(CST).strftime("%Y%m%d_%H%M%S")
    try:
        dt = datetime.fromisoformat(row["started_at"])
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        # 统一用北京时间命名，与库存的“当日去重”口径一致（IS-20）
        return dt.astimezone(CST).strftime("%Y%m%d_%H%M%S")
    except ValueError:
        return datetime.now(CST).strftime("%Y%m%d_%H%M%S")


def _dump_sheet(wb, title: str, headers: list[str], rows: Iterable[tuple]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    rows = list(rows)
    for row in rows:
        ws.append(list(row))
    from openpyxl.utils import get_column_letter

    all_rows = [headers] + [list(r) for r in rows]
    if not all_rows:
        return
    for i, col in enumerate(zip(*all_rows), start=1):
        w = max(len(str(v or "")) for v in col) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 12), 60)


def _sheet_榜单(wb, db: Database, round_id: int) -> None:
    rows = db.conn.execute(
        "SELECT shop_name, rank, offer_id, product_url, list_title FROM shop_offers "
        "WHERE round_id=? ORDER BY shop_key, rank",
        (round_id,),
    ).fetchall()
    _dump_sheet(wb, "店铺榜单",
                ["店铺", "排名", "offer_id", "商品链接", "列表标题"],
                [(r["shop_name"], r["rank"], r["offer_id"], r["product_url"], r["list_title"] or "")
                 for r in rows])


def _prev_success_map(db: Database, round_id: int) -> dict:
    prev = db.previous_complete_round(round_id)
    if prev is None:
        return {}
    return {_key(r): r for r in db.success_rows(int(prev["id"]))}


def _sheet_差分(wb, db: Database, round_id: int) -> None:
    prev_map = _prev_success_map(db, round_id)
    cur = db.success_rows(round_id)
    out = []
    for r in cur:
        p = prev_map.get(_key(r))
        out.append((
            r["shop_name"], r["product_url"], r["product_name"], r["sku_name"],
            r["sku_price"], (p["sku_price"] if p else None), r["sku_stock"],
            (p["sku_stock"] if p else None), r["stock_delta"],
            "补货/新增" if (r["stock_delta"] or 0) > 0 else "",
        ))
    _dump_sheet(wb, "SKU差分",
                ["店铺", "商品链接", "商品名称", "SKU名称", "本期价格", "上期价格",
                 "本期库存", "上期库存", "变化", "备注"],
                out)


def _sheet_汇总(wb, db: Database, round_id: int) -> None:
    prev_map = _prev_success_map(db, round_id)
    agg: dict[tuple, dict] = {}
    for r in db.success_rows(round_id):
        k = (r["shop_name"], r["shop_key"], r["product_url"], r["product_name"])
        a = agg.setdefault(k, {"cur_stock": 0, "prev_stock": 0, "skus": 0, "restock": 0})
        p = prev_map.get(_key(r))
        a["skus"] += 1
        a["cur_stock"] += r["sku_stock"] or 0
        if p is not None:
            a["prev_stock"] += p["sku_stock"] or 0
        if (r["stock_delta"] or 0) > 0:
            a["restock"] += 1
    out = []
    for (shop, _key2, url, name), a in agg.items():
        delta = a["cur_stock"] - a["prev_stock"]
        est_sales = max(a["prev_stock"] - a["cur_stock"], 0)
        out.append((shop, url, name, a["skus"], a["prev_stock"], a["cur_stock"],
                    delta, est_sales, a["restock"]))
    out.sort(key=lambda x: x[7], reverse=True)
    _dump_sheet(wb, "商品汇总",
                ["店铺", "商品链接", "商品名称", "SKU数", "上期总库存", "本期总库存",
                 "净变化", "估算销量", "补货SKU数"],
                out)


def _sheet_补货(wb, db: Database, round_id: int) -> None:
    prev_map = _prev_success_map(db, round_id)
    out = []
    for r in db.success_rows(round_id):
        if (r["stock_delta"] or 0) > 0 or _key(r) not in prev_map:
            out.append((r["shop_name"], r["product_url"], r["product_name"], r["sku_name"],
                        r["sku_stock"], r["stock_delta"]))
    _dump_sheet(wb, "疑似补货",
                ["店铺", "商品链接", "商品名称", "SKU名称", "本期库存", "变化"], out)


def _sheet_榜单变化(wb, db: Database, round_id: int) -> None:
    cur = db.conn.execute(
        "SELECT shop_key, offer_id, rank FROM shop_offers WHERE round_id=?",
        (round_id,),
    ).fetchall()
    prev = db.previous_complete_round(round_id)
    prev_rows = []
    if prev is not None:
        prev_rows = db.conn.execute(
            "SELECT shop_key, offer_id, rank FROM shop_offers WHERE round_id=?",
            (int(prev["id"]),),
        ).fetchall()
    cur_map = {(r["shop_key"], r["offer_id"]): r["rank"] for r in cur}
    prev_map = {(r["shop_key"], r["offer_id"]): r["rank"] for r in prev_rows}
    out = []
    for key, rank in cur_map.items():
        old = prev_map.get(key)
        out.append((key[0], key[1], "新上榜" if old is None else f"{old}→{rank}",
                    "新上榜" if old is None else ""))
    for key, old in prev_map.items():
        if key not in cur_map:
            out.append((key[0], key[1], f"原第{old}名", "今日不在榜"))
    out.sort(key=lambda x: (x[0], x[2]))
    _dump_sheet(wb, "榜单变化", ["店铺", "offer_id", "排名变化", "状态"], out)


def _sheet_失败(wb, db: Database, round_id: int) -> None:
    rows = db.failed_rows(round_id)
    _dump_sheet(wb, "失败清单",
                ["店铺", "offer_id", "商品链接", "状态", "尝试次数", "说明"],
                [(r["shop_name"], r["offer_id"], r["product_url"], r["page_status"],
                  r["attempt"], r["detail_note"] or "") for r in rows])
