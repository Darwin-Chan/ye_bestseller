import tempfile
import unittest
from pathlib import Path

from bestseller_monitor.config import Config
from bestseller_monitor.db import Database, connect
from bestseller_monitor.report import export_round_excel, update_stock_deltas

REPO = Path(__file__).resolve().parents[1]


def _row(rid, shop_key, offer_id, sku_id, stock, name="商品", attempt=1, delta=None):
    return {
        "round_id": rid, "shop_key": shop_key,
        "shop_url": f"https://{shop_key}.example/", "shop_name": f"店铺{shop_key}",
        "offer_id": offer_id, "product_url": f"https://detail.1688.com/offer/{offer_id}.html",
        "product_name": name, "sku_id": sku_id, "sku_name": f"SKU-{sku_id}",
        "sku_price": 1.0, "sku_stock": stock, "stock_delta": delta,
        "collected_at": "2026-09-04T00:00:00+00:00",
        "page_status": "成功", "attempt": attempt,
    }


class ReportTests(unittest.TestCase):
    def test_update_stock_deltas(self):
        with tempfile.TemporaryDirectory() as tmp:
            conn = connect(Path(tmp) / "t.db")
            db = Database(conn)
            r1 = db.start_or_resume()
            db.finish_round(r1)
            db.save_snapshot_rows(r1, "A", [_row(r1, "A", "111", "s1", 200)])

            r2 = db.start_or_resume()
            db.save_snapshot_rows(r2, "A", [_row(r2, "A", "111", "s1", 160)])
            db.finish_round(r2)
            update_stock_deltas(db, r2)
            rows = db.success_rows(r2)
            self.assertEqual(rows[0]["stock_delta"], -40)
            conn.close()

    def test_excel_report_builds(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl 未安装")
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            cfg = (
                Config.from_file(REPO / "config" / "config.toml", root=REPO)
                .replace(
                    db_file=base / "b.db",
                    data_dir=base / "data",
                    output_dir=base / "output",
                    logs_dir=base / "logs",
                    screenshot_dir=base / "shots",
                    raw_page_dir=base / "raw",
                    profile_dir=base / "profile",
                )
            )
            cfg.ensure_dirs()
            conn = connect(cfg.db_file)
            db = Database(conn)

            r1 = db.start_or_resume()
            db.save_shop_offers(
                r1, "A", "https://a.example/", "店铺A",
                [(1, "111", "https://detail.1688.com/offer/111.html", "商品1", "")],
                1,
            )
            db.save_snapshot_rows(r1, "A", [_row(r1, "A", "111", "s1", 200)])
            db.finish_round(r1)

            r2 = db.start_or_resume()
            db.save_shop_offers(
                r2, "A", "https://a.example/", "店铺A",
                [(1, "111", "https://detail.1688.com/offer/111.html", "商品1", ""),
                 (2, "222", "https://detail.1688.com/offer/222.html", "商品2", "")],
                1,
            )
            db.save_snapshot_rows(r2, "A", [_row(r2, "A", "111", "s1", 160)])
            db.finish_round(r2)

            update_stock_deltas(db, r2)
            xlsx = export_round_excel(db, cfg, r2)
            self.assertTrue(xlsx.exists())
            wb = openpyxl.load_workbook(xlsx)
            self.assertIn("SKU差分", wb.sheetnames)
            self.assertIn("榜单变化", wb.sheetnames)
            wb.close()
            conn.close()


if __name__ == "__main__":
    unittest.main()
